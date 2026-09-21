"""Genera archivos .xlsx minimos con el mismo acomodo que los de planta.

Asi los tests no dependen de los documentos reales de Florence, que no se
versionan.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from hiloflo.parametros import COLUMNAS, HOJA_PREDETERMINADA


def crear_wi(ruta: Path, velocidades: dict[float, dict[int, float]]) -> Path:
    """Arma un WI de juguete: {diametro: {columna: mm/s}}.

    Replica el acomodo real: encabezados de bloque arriba y la columna A con
    el diametro de alambre estirado.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_PREDETERMINADA

    ws.cell(1, 1, "ITW Line Speed\n[mm/s]")
    ws.cell(2, 1, "Drawn Wire Rod Ø\n[mm]")
    for columna, (lineas, *_) in COLUMNAS.items():
        ws.cell(2, columna, "/".join(lineas))

    fila = 3
    for diametro, por_columna in sorted(velocidades.items()):
        ws.cell(fila, 1, diametro)
        for columna, mm_s in por_columna.items():
            ws.cell(fila, columna, mm_s)
        fila += 1

    wb.save(ruta)
    return ruta


ENCABEZADOS_SCHEDULE = [
    "Work Center",
    "Material Number",
    "Material Description",
    "Important Notes",
    "Order",
    "PO Printed Yes/No",
    "Drawn Yes/No",
    "Operation Quantity (MEINH)",
    "Customer PO",
]


def crear_schedule(ruta: Path, renglones: list[tuple]) -> Path:
    """Arma un schedule de juguete.

    Cada renglon es (work_center, material, descripcion, notas, orden, kg).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.cell(1, 1, "CSW Production Schedule (prueba)")
    for col, encabezado in enumerate(ENCABEZADOS_SCHEDULE, start=1):
        ws.cell(2, col, encabezado)

    for fila, (wc, material, descripcion, notas, orden, kg) in enumerate(
        renglones, start=3
    ):
        ws.cell(fila, 1, wc)
        ws.cell(fila, 2, material)
        ws.cell(fila, 3, descripcion)
        ws.cell(fila, 4, notas)
        ws.cell(fila, 5, orden)
        ws.cell(fila, 8, kg)

    wb.save(ruta)
    return ruta
