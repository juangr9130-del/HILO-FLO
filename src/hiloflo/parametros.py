"""Lector del WI-FLO-CSW-P-526 "ITW Process Parameters".

De ese documento solo nos interesa la tabla "ITW Line Speed [mm/s]": el
diametro de alambre estirado en la columna A y la velocidad de cada linea
en las columnas C..P.

La tabla viene partida en varios bloques (por el salto de pagina del
documento) pero el acomodo de columnas se repite identico en todos, asi
que se recorre la hoja completa y se toma cualquier renglon cuya columna A
sea numerica.

Acomodo de columnas (confirmado en los 8 bloques del documento):

    C   ITW-2   devanador Neturen
    D   ITW-2   devanador DEM, grado 9254
    E   ITW-2   devanador DEM, grado 1065
    F   ITW-3
    G   ITW-5 y ITW-6
    H   ITW-1
    I   ITW-7, ITW-8 y ITW-9
    J   ITW-10  NON SLM
    K   ITW-10  SLM
    L   ITW-4
    M   ITW-11
    N   ITW-12
    O   ITW-13
    P   ITW-14
"""

from __future__ import annotations

from pathlib import Path

from .modelos import PuntoVelocidad

HOJA_PREDETERMINADA = "Anlagen - Setup "

# columna -> (lineas, winder, grado, slm)
COLUMNAS: dict[int, tuple[tuple[str, ...], str | None, str | None, bool | None]] = {
    3: (("ITW-2",), "NETUREN", None, None),
    4: (("ITW-2",), "DEM", "9254", None),
    5: (("ITW-2",), "DEM", "1065", None),
    6: (("ITW-3",), None, None, None),
    7: (("ITW-5", "ITW-6"), None, None, None),
    8: (("ITW-1",), None, None, None),
    9: (("ITW-7", "ITW-8", "ITW-9"), None, None, None),
    10: (("ITW-10",), None, None, False),
    11: (("ITW-10",), None, None, True),
    12: (("ITW-4",), None, None, None),
    13: (("ITW-11",), None, None, None),
    14: (("ITW-12",), None, None, None),
    15: (("ITW-13",), None, None, None),
    16: (("ITW-14",), None, None, None),
}

# El diametro de la columna A trae ruido de punto flotante (6.25000000000001).
DECIMALES_DIAMETRO = 2

# Rango plausible de diametro de alambre estirado en las lineas ITW (mm).
DIAMETRO_MIN = 4.0
DIAMETRO_MAX = 30.0


def leer_velocidades(
    ruta: str | Path, hoja: str | None = None
) -> list[PuntoVelocidad]:
    """Extrae todos los puntos de la tabla ITW Line Speed."""
    ruta = Path(ruta)
    if not ruta.exists():
        raise FileNotFoundError(f"no existe el archivo: {ruta}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise ImportError(
            "se requiere openpyxl para leer el WI "
            "(pip install -r requirements.txt)"
        ) from exc

    wb = load_workbook(ruta, data_only=True)
    ws = _elegir_hoja(wb, hoja)

    puntos: list[PuntoVelocidad] = []
    for fila in range(1, ws.max_row + 1):
        diametro = ws.cell(fila, 1).value
        if not isinstance(diametro, (int, float)):
            continue
        diametro = round(float(diametro), DECIMALES_DIAMETRO)
        if not DIAMETRO_MIN <= diametro <= DIAMETRO_MAX:
            continue
        for columna, (lineas, winder, grado, slm) in COLUMNAS.items():
            velocidad = ws.cell(fila, columna).value
            if not isinstance(velocidad, (int, float)) or velocidad <= 0:
                continue  # celda vacia = esa linea no corre ese diametro
            for linea in lineas:
                puntos.append(
                    PuntoVelocidad(
                        linea=linea,
                        diametro_mm=diametro,
                        mm_s=float(velocidad),
                        winder=winder,
                        grado=grado,
                        slm=slm,
                    )
                )
    if not puntos:
        raise ValueError(
            f"no se encontro ninguna velocidad en {ruta.name}; "
            "revisa que la hoja sea la del ITW Line Speed"
        )
    return puntos


def _elegir_hoja(wb, hoja: str | None):
    if hoja is not None:
        return wb[hoja]
    if HOJA_PREDETERMINADA in wb.sheetnames:
        return wb[HOJA_PREDETERMINADA]
    # El nombre de la hoja trae un espacio al final en el documento original;
    # se compara sin espacios por si lo corrigen.
    objetivo = HOJA_PREDETERMINADA.strip().lower()
    for nombre in wb.sheetnames:
        if nombre.strip().lower() == objetivo:
            return wb[nombre]
    return wb.worksheets[0]
